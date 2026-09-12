import json
from pathlib import Path
from openpyxl import load_workbook

root = Path('/home/ubuntu/tradebilia-isolated-development')
data = json.loads((root / 'scripts/test_ai_query_audit.json').read_text())
rows = [row for row in data['rows'] if row['category'] == 'sports_cards']
assert len(rows) == 4, len(rows)
by_type = {row['itemType']: row for row in rows}
assert set(by_type) == {'single_card', 'card_set', 'unopened_product', 'collection_lot'}
assert 'Player' in by_type['single_card']['query']['activeEbay']
assert 'Player' not in by_type['card_set']['query']['activeEbay']
assert 'Product Format' in by_type['unopened_product']['query']['activeEbay']
assert 'Sport' in by_type['unopened_product']['query']['activeEbay']
assert 'Product Name' in by_type['unopened_product']['query']['excludedFields']
assert 'Condition' in by_type['unopened_product']['query']['excludedFields']
assert 'FASC' in by_type['unopened_product']['query']['conditionalRules']
assert 'Collection-lot fields' in by_type['collection_lot']['query']['notes']

workbook = root / 'exports' / 'tradebilia_test_ai_query_audit.xlsx'
wb = load_workbook(workbook, read_only=False, data_only=True)
assert set(['Read Me', 'Category Item Types', 'Query Mapping', 'Field Detail', 'Provider Behavior', 'Upload Template']).issubset(wb.sheetnames)
ws = wb['Query Mapping']
headers = [cell.value for cell in ws[1]]
assert headers[2] == 'Active eBay Query Fields'
query_rows = {ws.cell(row, 1).value + ':' + ws.cell(row, 2).value: [ws.cell(row, col).value for col in range(1, 9)] for row in range(2, ws.max_row + 1)}
assert 'sports_cards:single_card' in query_rows
assert 'sports_cards:unopened_product' in query_rows
assert query_rows['sports_cards:single_card'][2] != query_rows['sports_cards:unopened_product'][2]
print(f'Validated {len(data["rows"])} mappings, {len(wb["Field Detail"]["A"])} field rows, and distinct Sports Cards query behavior.')
